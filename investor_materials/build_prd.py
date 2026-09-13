from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / 'investor_materials' / 'Hirewise_Product_Requirements_Document.docx'
LOGO = ROOT / 'static' / 'logo.png'

INK = '0D3932'
GREEN = '1D695B'
LIME = 'D5EB8B'
PALE = 'F1F5E8'
MUTED = RGBColor(85, 105, 97)


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    element = tc_pr.find(qn('w:shd'))
    if element is None:
        element = OxmlElement('w:shd')
        tc_pr.append(element)
    element.set(qn('w:fill'), fill)


def set_cell_text(cell, text, bold=False, color=INK, size=8.5):
    cell.text = ''
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.space_before = Pt(0)
    run = p.add_run(text)
    run.bold = bold
    run.font.name = 'Aptos'
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_heading(doc, text, level=1):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4 if level == 1 else 2)
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run(text)
    run.bold = True
    run.font.name = 'Aptos Display'
    run.font.size = Pt(14 if level == 1 else 10)
    run.font.color.rgb = RGBColor.from_string(INK)
    return p


def add_body(doc, text, size=8.5, bold_prefix=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2.5)
    p.paragraph_format.line_spacing = 1.0
    if bold_prefix and text.startswith(bold_prefix):
        first = p.add_run(bold_prefix)
        first.bold = True
        first.font.color.rgb = RGBColor.from_string(INK)
        rest = p.add_run(text[len(bold_prefix):])
        rest.font.color.rgb = MUTED
        runs = (first, rest)
    else:
        run = p.add_run(text)
        run.font.color.rgb = MUTED
        runs = (run,)
    for run in runs:
        run.font.name = 'Aptos'
        run.font.size = Pt(size)
    return p


def add_bullets(doc, items, size=8.2):
    for item in items:
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.left_indent = Inches(.16)
        p.paragraph_format.first_line_indent = Inches(-.12)
        p.paragraph_format.space_after = Pt(1.5)
        p.paragraph_format.line_spacing = 1.0
        run = p.add_run(item)
        run.font.name = 'Aptos'
        run.font.size = Pt(size)
        run.font.color.rgb = MUTED


def add_header(doc, page_label):
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.columns[0].width = Inches(5.65)
    table.columns[1].width = Inches(1.15)
    table.cell(0, 0).paragraphs[0].add_run().add_picture(str(LOGO), width=Inches(.72))
    p = table.cell(0, 1).paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = p.add_run(page_label)
    run.bold = True
    run.font.name = 'Aptos'
    run.font.size = Pt(7.5)
    run.font.color.rgb = MUTED
    for cell in table.rows[0].cells:
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


doc = Document()
section = doc.sections[0]
section.top_margin = Inches(.42)
section.bottom_margin = Inches(.42)
section.left_margin = Inches(.55)
section.right_margin = Inches(.55)
section.header_distance = Inches(.2)
section.footer_distance = Inches(.2)

styles = doc.styles
styles['Normal'].font.name = 'Aptos'
styles['Normal'].font.size = Pt(8.5)
styles['Normal'].paragraph_format.space_after = Pt(2)

add_header(doc, 'PRODUCT REQUIREMENTS DOCUMENT · 01/02')

p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(3)
r = p.add_run('Hirewise — Product Requirements Document')
r.bold = True
r.font.name = 'Aptos Display'
r.font.size = Pt(21)
r.font.color.rgb = RGBColor.from_string(INK)
add_body(doc, 'Version 1.0 · 13 September 2026 · Owner: Hirewise · Status: MVP requirements baseline', 8)

summary = doc.add_table(rows=2, cols=4)
summary.alignment = WD_TABLE_ALIGNMENT.CENTER
summary.autofit = False
labels = [('Market', 'Qatar first'), ('Platform', 'B2B recruitment SaaS'), ('Users', 'Candidates, recruiters, admins'), ('Decision model', 'AI-assisted, human-controlled')]
for i, (label, value) in enumerate(labels):
    set_cell_text(summary.cell(0, i), label.upper(), True, GREEN, 7)
    set_cell_text(summary.cell(1, i), value, True, INK, 8)
    shade(summary.cell(0, i), LIME)
    shade(summary.cell(1, i), PALE)

add_heading(doc, '1. Product vision and problem')
add_body(doc, 'Hirewise is a privacy-first recruitment platform that connects job requirements, applications, evidence-based screening, text interviews and human decisions in one auditable workflow. It targets Qatar initially while preserving a path to multi-country operation.')
add_body(doc, 'Problem: Recruiters manage high application volume across fragmented tools, while candidates have limited clarity about data use and AI-assisted evaluation. Keyword filters and opaque scores can hide missing evidence, weaken accountability and create poor candidate experiences.', bold_prefix='Problem:')

add_heading(doc, '2. Objectives and success criteria')
add_bullets(doc, [
    'Enable a recruiter to register, create/verify a company, publish a vacancy, receive applications, review evidence, request an interview, record a decision and export a report.',
    'Enable a candidate to find a job, apply with or without an account, upload a CV privately, receive a reference, complete an invited text interview, track or withdraw an application and exercise privacy rights.',
    'Give administrators MFA-protected controls for companies, jobs, pricing, privacy requests, configuration, security events and audit history.',
    'Preserve every significant screening, interview and status event so an authorized reviewer can explain why a candidate advanced or did not advance.',
])

table = doc.add_table(rows=1, cols=3)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
for i, value in enumerate(['Pilot KPI', 'Target hypothesis', 'Measurement']):
    set_cell_text(table.cell(0, i), value, True, 'FFFFFF', 7.5)
    shade(table.cell(0, i), INK)
for values in [
    ('Application completion', '≥70%', 'Started vs. submitted'),
    ('Time to first review', '≤2 business days', 'Received to recruiter view'),
    ('Screening explainability', '100% of scores', 'Criteria with evidence/unknown label'),
    ('Tenant isolation', '0 unauthorized disclosures', 'Automated tests + security review'),
]:
    row = table.add_row().cells
    for i, value in enumerate(values):
        set_cell_text(row[i], value, i == 0, INK if i == 0 else GREEN, 7.7)
        shade(row[i], PALE if len(table.rows) % 2 == 0 else 'FFFFFF')

add_heading(doc, '3. Users and permissions')
add_body(doc, 'Candidate: own profile, documents, applications, interviews, notifications and privacy requests. Guest: vacancy-scoped application and expiring email-verified access. Recruiter: active-company jobs and verified applicants only; never another tenant. Administrator: MFA-protected platform management with logged actions. Internal reviewer is future scope with assignment-only access.')

add_heading(doc, '4. MVP functional requirements')
add_bullets(doc, [
    'Authentication: email verification, secure login/reset, Argon2 passwords, database sessions, administrator TOTP MFA and server-side RBAC.',
    'Company/jobs: company membership and approval; draft, moderation, publish, pause, close, archive and duplicate; structured job-related criteria with versioned weights.',
    'Applications: guest or account submission, unique reference, consent version, screening answers, private CV and append-only stage timeline.',
    'Documents: PDF/DOCX/TXT validation, 5 MB limit, private quarantine, malware scan, bounded extraction, integrity hash and short-lived authorized downloads.',
])

doc.add_page_break()

add_header(doc, 'PRODUCT REQUIREMENTS DOCUMENT · 02/02')
add_heading(doc, '5. Matching, interview and commercial requirements')
add_bullets(doc, [
    'Matching combines objective mandatory rules, structured scoring, reviewed skill equivalence and contextual/semantic comparison. Every result stores policy/model/prompt version, sources, score, explanation and unknowns. Protected or irrelevant attributes are excluded.',
    'AI output is advisory. Default categories are Strong Match, Potential Match, Review Required and Low Match. No irreversible employment decision may be automated. Human override preserves both recommendations and requires a reason when configured.',
    'Text interview questions derive from recruiter-approved job competencies. Candidates receive disclosure and consent before scheduling. Reports show questions, answers, rubric, evidence, confidence and required human review. No facial, emotion, accent or appearance analysis.',
    'Initial commercial model is pay per vacancy with configurable QAR/USD packages. Payment uses a replaceable provider interface, hosted checkout, server verification, signed/idempotent webhooks and permission-controlled invoices. No card number or CVV is stored.',
])

add_heading(doc, '6. Core workflow and application states')
workflow = doc.add_table(rows=1, cols=5)
workflow.alignment = WD_TABLE_ALIGNMENT.CENTER
for i, value in enumerate(['1 · Define', '2 · Apply', '3 · Screen', '4 · Interview', '5 · Decide']):
    set_cell_text(workflow.cell(0, i), value, True, INK, 8)
    shade(workflow.cell(0, i), LIME if i in (0, 4) else PALE)
add_body(doc, 'Application stages: Received → Under Review → AI Screening → Interview Requested/Scheduled/Completed → Human Review → Shortlisted → Employer Interview → Offer → Hired; terminal alternatives: Rejected, Withdrawn or Archived. Every transition records actor, UTC timestamp, previous/new state and reason where required.', 8)

add_heading(doc, '7. Privacy, security and non-functional requirements')
security = doc.add_table(rows=1, cols=2)
security.alignment = WD_TABLE_ALIGNMENT.CENTER
requirements = [
    ('Privacy', 'Data minimization; purpose-specific notices/consent; no scraping or brokers; configurable retention/legal hold; access, correction, export, withdrawal and deletion workflows.'),
    ('Authorization', 'Deny by default; server-side role and tenant checks; candidate/recruiter projections; private documents; audit all sensitive access and exports.'),
    ('Application security', 'TLS/HSTS, secure cookies, CSRF, CSP, input validation, escaped output, throttling, secret management, dependency scanning and least privilege.'),
    ('Availability', 'AI/payment/email failure never loses an application. Browsing and human workflows continue when AI is unavailable. Background tasks retry with bounded leases.'),
    ('Performance', 'Paginated dashboards; no thousands of candidates in browser memory; indexed tenant/status/time queries; expensive parsing/matching/reporting occurs asynchronously.'),
    ('Operations', 'PostgreSQL, private object storage, workers, structured logs without CV text, monitoring, encrypted backups, restoration tests, rollback and separate dev/staging/production.'),
]
for i, (label, body) in enumerate(requirements):
    if i and i % 2 == 0:
        security.add_row()
    row = security.rows[i // 2]
    cell = row.cells[i % 2]
    cell.text = ''
    shade(cell, PALE if i % 3 else 'FFFFFF')
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(1)
    r = p.add_run(label + '\n')
    r.bold = True
    r.font.name = 'Aptos'
    r.font.size = Pt(8.5)
    r.font.color.rgb = RGBColor.from_string(INK)
    r = p.add_run(body)
    r.font.name = 'Aptos'
    r.font.size = Pt(7.4)
    r.font.color.rgb = MUTED

add_heading(doc, '8. Release scope and dependencies')
scope = doc.add_table(rows=1, cols=2)
scope.alignment = WD_TABLE_ALIGNMENT.CENTER
for i, value in enumerate(['MVP release includes', 'Excluded / gated']):
    set_cell_text(scope.cell(0, i), value, True, 'FFFFFF', 8)
    shade(scope.cell(0, i), INK)
left = 'Accounts and RBAC; company/jobs; guest and candidate applications; private uploads; evidence-linked matching; text interviews; sandbox payments; audit, privacy and reports.'
right = 'Audio/video and emotion analysis; social scraping; autonomous rejection; live AI/payment activation without approved provider; public intake before hosting, legal, security and recovery sign-off.'
row = scope.add_row().cells
set_cell_text(row[0], left, False, INK, 7.7)
shade(row[0], PALE)
set_cell_text(row[1], right, False, INK, 7.7)
shade(row[1], PALE)

add_heading(doc, '9. Acceptance gates and current status')
add_bullets(doc, [
    'Critical authorization, upload, prompt-injection, payment-tampering, session and report-export tests pass against PostgreSQL in CI; deployed storage/scanner and penetration testing are additionally required.',
    'Production requires approved hosting/data regions, private storage, SMTP, malware scanning, monitoring, backups with successful restore evidence, qualified Qatar legal/accounting review and a regulated/acquirer-compatible payment agreement.',
    'Current repository contains a working local baseline and internal payment/evidence engines. Full semantic AI, live payment integration and production infrastructure remain launch dependencies and must not be represented as completed.',
])

add_heading(doc, '10. Product decisions and assumptions')
add_body(doc, 'Modular Django monolith, PostgreSQL, S3-compatible private storage and background workers are the initial architecture. Qatar, QAR and Asia/Qatar are defaults; English and text interviews launch first; Arabic/RTL and multi-country configuration follow. Hirewise is the working brand. Pricing, retention periods, scoring policies and investor targets require owner and professional validation before production.')

for section in doc.sections:
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run('HIREWISE  ·  CONFIDENTIAL  ·  PRODUCT REQUIREMENTS BASELINE')
    run.font.name = 'Aptos'
    run.font.size = Pt(6.5)
    run.font.color.rgb = MUTED

doc.core_properties.title = 'Hirewise Product Requirements Document'
doc.core_properties.subject = 'Two-page product requirements baseline'
doc.core_properties.author = 'Hirewise'
doc.core_properties.comments = 'Requirements baseline; legal, commercial and production dependencies remain subject to approval.'
doc.save(OUTPUT)
print(f'Created {OUTPUT}')
