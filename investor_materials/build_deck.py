from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / 'investor_materials' / 'Hirewise_Investor_Deck.pptx'
LOGO = ROOT / 'static' / 'logo.png'

INK = RGBColor(13, 57, 50)
INK_2 = RGBColor(26, 78, 68)
LIME = RGBColor(213, 235, 139)
CREAM = RGBColor(248, 248, 239)
WHITE = RGBColor(255, 255, 255)
MUTED = RGBColor(95, 113, 105)
PALE = RGBColor(235, 241, 223)
LINE = RGBColor(218, 226, 214)
BLUE = RGBColor(19, 118, 183)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)


def add_bg(slide, color=WHITE):
    bg = slide.background.fill
    bg.solid()
    bg.fore_color.rgb = color


def add_text(slide, text, x, y, w, h, size=18, color=INK, bold=False,
             font='Aptos', align=PP_ALIGN.LEFT, margin=0, valign=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = Inches(margin)
    frame.vertical_anchor = valign
    p = frame.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return box


def add_rich_lines(slide, lines, x, y, w, h, size=18, color=INK, bullet=False, spacing=10):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0)
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.font.name = 'Aptos'
        p.font.size = Pt(size)
        p.font.color.rgb = color
        p.space_after = Pt(spacing)
        p.level = 0
        if bullet:
            p.text = '•  ' + line
    return box


def rect(slide, x, y, w, h, fill=WHITE, line=LINE, radius=True):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line
    shape.line.width = Pt(1)
    return shape


def circle(slide, x, y, d, fill=LIME, line=LIME):
    shape = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line
    return shape


def header(slide, number, eyebrow, title, subtitle=None, dark=False):
    color = WHITE if dark else INK
    muted = PALE if dark else MUTED
    add_text(slide, f'{number:02d}', 0.65, 0.42, 0.5, 0.25, 9, LIME if dark else MUTED, True)
    add_text(slide, eyebrow.upper(), 1.15, 0.42, 4.8, 0.25, 9, LIME if dark else MUTED, True)
    add_text(slide, title, 0.65, 0.88, 11.9, 0.95, 31, color, True)
    if subtitle:
        add_text(slide, subtitle, 0.67, 1.76, 11.6, 0.48, 13, muted)


def footer(slide, label='CONFIDENTIAL · INVESTOR DISCUSSION'):
    add_text(slide, label, 0.65, 7.15, 5, 0.18, 7, MUTED, True)
    add_text(slide, 'HIREWISE  ·  SEPTEMBER 2026', 9.9, 7.15, 2.8, 0.18, 7, MUTED, True, align=PP_ALIGN.RIGHT)


def logo(slide, x=11.55, y=0.32, w=1.1):
    slide.shapes.add_picture(str(LOGO), Inches(x), Inches(y), width=Inches(w), height=Inches(w))


def new_slide(bg=WHITE):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide, bg)
    return slide


# 1 — Cover
s = new_slide(INK)
circle(s, 8.2, -1.6, 6.3, INK_2, INK_2)
circle(s, 9.0, 0.0, 4.4, LIME, LIME)
circle(s, 9.55, 0.55, 3.3, CREAM, CREAM)
s.shapes.add_picture(str(LOGO), Inches(9.9), Inches(0.9), width=Inches(2.6), height=Inches(2.6))
add_text(s, 'HIREWISE', 0.8, 0.72, 4, 0.25, 10, LIME, True)
add_text(s, 'Smarter hiring.\nFairer decisions.', 0.8, 1.35, 7.1, 1.8, 38, WHITE, True)
add_text(s, 'AI-assisted recruitment infrastructure for Qatar—\nbuilt around evidence, privacy and human judgment.', 0.83, 3.45, 6.6, 1.25, 16, PALE)
rect(s, 0.8, 5.15, 5.75, 1.08, INK_2, INK_2)
add_text(s, 'INVESTOR DISCUSSION', 1.1, 5.45, 2.6, 0.25, 9, LIME, True)
add_text(s, 'Pre-seed concept & MVP status', 3.55, 5.4, 2.65, 0.35, 13, WHITE)
add_text(s, 'September 2026  ·  Doha, Qatar', 0.83, 6.85, 5, 0.24, 9, PALE)

# 2 — Problem
s = new_slide()
header(s, 2, 'The problem', 'Hiring teams see volume. They need evidence.', 'Recruitment tools often optimize workflow without improving the quality, transparency or privacy of evaluation.')
problems = [
    ('01', 'Fragmented evidence', 'CVs, screening answers, notes and interviews live in disconnected systems.'),
    ('02', 'Opaque screening', 'Keyword filters and black-box scores make decisions difficult to explain or audit.'),
    ('03', 'Candidate trust gap', 'Applicants have limited visibility into data use, retention and the role of AI.'),
]
for i, (num, title, body) in enumerate(problems):
    x = 0.67 + i * 4.18
    rect(s, x, 2.45, 3.8, 3.55, CREAM)
    add_text(s, num, x + .28, 2.73, .45, .3, 11, MUTED, True)
    circle(s, x + 2.95, 2.62, .5, LIME, LIME)
    add_text(s, title, x + .28, 3.38, 3.1, .48, 20, INK, True)
    add_text(s, body, x + .28, 4.08, 3.05, 1.35, 11, MUTED)
    add_text(s, '→', x + .28, 5.42, .5, .3, 17, INK, True)
footer(s)

# 3 — Solution
s = new_slide(INK)
header(s, 3, 'The solution', 'One accountable hiring journey.', 'Hirewise connects vacancy requirements, candidate evidence, screening rationale, interviews and human decisions.', dark=True)
steps = [('1', 'Define', 'Structured, job-related criteria'), ('2', 'Apply', 'Private guest or account application'), ('3', 'Understand', 'Evidence-linked matching & gaps'), ('4', 'Interview', 'Approved questions and rubric'), ('5', 'Decide', 'Human action with reason & history')]
for i, (num, title, body) in enumerate(steps):
    x = .65 + i * 2.5
    circle(s, x, 2.55, .58, LIME, LIME)
    add_text(s, num, x, 2.67, .58, .22, 11, INK, True, align=PP_ALIGN.CENTER)
    if i < 4:
        line = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x + .68), Inches(2.81), Inches(1.72), Inches(.018))
        line.fill.solid()
        line.fill.fore_color.rgb = RGBColor(89, 126, 116)
        line.line.fill.background()
    add_text(s, title, x, 3.42, 2.1, .36, 17, WHITE, True)
    add_text(s, body, x, 3.96, 2.08, 1.05, 10, PALE)
rect(s, .65, 5.42, 12.0, .88, INK_2, INK_2)
add_text(s, 'AI assists. Evidence explains. People decide.', 1.0, 5.69, 8.4, .32, 19, WHITE, True)
add_text(s, 'NO AUTONOMOUS EMPLOYMENT DECISIONS', 9.45, 5.73, 2.8, .2, 8, LIME, True, align=PP_ALIGN.RIGHT)
footer(s)

# 4 — Product
s = new_slide()
header(s, 4, 'Product', 'A modular recruitment operating system.', 'Four experiences share one secure, traceable data model.')
modules = [
    ('Candidate', 'Find roles · guest apply · upload CV · track · interview · privacy controls'),
    ('Recruiter', 'Companies · vacancies · applicants · evidence · pipeline · reports'),
    ('Administrator', 'Approvals · moderation · pricing · audit · privacy · configuration'),
    ('Intelligence layer', 'Document extraction · versioned matching · evidence · interview assessment'),
]
for i, (title, body) in enumerate(modules):
    x = .67 + (i % 2) * 6.12
    y = 2.35 + (i // 2) * 2.05
    rect(s, x, y, 5.72, 1.65, PALE if i == 3 else CREAM)
    add_text(s, title, x + .28, y + .25, 2.4, .35, 17, INK, True)
    add_text(s, body, x + .28, y + .78, 5.05, .62, 10, MUTED)
    add_text(s, '↗', x + 5.02, y + .22, .35, .3, 17, INK, True)
add_text(s, 'Provider-independent architecture', .7, 6.48, 3.0, .28, 11, INK, True)
add_text(s, 'PostgreSQL · private object storage · background workers · replaceable AI/payment adapters', 3.8, 6.49, 8.6, .3, 11, MUTED)
footer(s)

# 5 — Differentiation
s = new_slide(CREAM)
header(s, 5, 'Why Hirewise', 'Trust is product infrastructure.', 'Designed for the realities of recruitment data and AI-assisted evaluation.')
items = [
    ('Evidence, not a mystery score', 'Every criterion points to a source—or clearly says evidence is unknown.'),
    ('Privacy from the first application', 'No secret scraping, no data brokers, no public CV directories.'),
    ('Human control by design', 'AI recommendations never become irreversible employment decisions.'),
    ('Traceable over time', 'Policy, model, prompt, score, status and override history remain versioned.'),
    ('Qatar first, portable later', 'QAR, Asia/Qatar, local payment abstraction and multi-country architecture.'),
    ('Company isolation', 'Recruiters access only the vacancies and applicants their active membership permits.'),
]
for i, (title, body) in enumerate(items):
    x = .72 + (i % 3) * 4.2
    y = 2.35 + (i // 3) * 2.15
    circle(s, x, y + .02, .38, LIME, LIME)
    add_text(s, '✓', x, y + .09, .38, .15, 9, INK, True, align=PP_ALIGN.CENTER)
    add_text(s, title, x + .58, y, 3.25, .38, 15, INK, True)
    add_text(s, body, x + .58, y + .58, 3.15, .92, 9, MUTED)
footer(s)

# 6 — Why now
s = new_slide()
header(s, 6, 'Why Qatar, why now', 'A national tailwind for digital talent.', 'Official strategy creates a timely entry point; Hirewise still needs customer discovery to quantify its serviceable market.')
rect(s, .7, 2.35, 4.0, 3.35, INK, INK)
add_text(s, '~QAR 40B', 1.05, 2.85, 3.3, .68, 31, LIME, True)
add_text(s, 'Digital Agenda 2030\napproximate cumulative impact', 1.05, 3.8, 3.2, .72, 14, WHITE, True)
add_text(s, 'Official MCIT strategy target', 1.05, 5.13, 2.6, .25, 9, PALE)
rect(s, 4.95, 2.35, 3.25, 3.35, PALE, PALE)
add_text(s, '26,000', 5.3, 2.85, 2.5, .68, 31, INK, True)
add_text(s, 'ICT jobs targeted\nby 2030', 5.3, 3.8, 2.3, .72, 14, INK, True)
add_text(s, 'Official MCIT strategy target', 5.3, 5.13, 2.4, .25, 9, MUTED)
rect(s, 8.45, 2.35, 4.15, 3.35, WHITE, LINE)
add_text(s, 'The opening', 8.8, 2.75, 2.5, .35, 18, INK, True)
add_rich_lines(s, ['Growing digital-talent demand', 'National focus on data & emerging tech', 'Need for trusted, efficient hiring infrastructure'], 8.8, 3.35, 3.35, 1.7, 13, MUTED, True, 8)
add_text(s, 'Source: Qatar MCIT Digital Agenda 2030 summary', .72, 6.43, 6.4, .25, 9, MUTED)
footer(s)

# 7 — Market wedge
s = new_slide(INK)
header(s, 7, 'Market strategy', 'Start narrow. Earn the right to expand.', 'A bottom-up commercial wedge avoids pretending that a broad global HR-tech market equals Hirewise’s obtainable market.', dark=True)
segments = [
    ('1', 'Beachhead', 'Qatar startups & SMEs hiring recurring professional and digital roles'),
    ('2', 'Expansion', 'Recruitment agencies and multi-seat mid-market employers'),
    ('3', 'Regional', 'GCC localization, enterprise controls, ATS/HRIS integrations'),
]
for i, (n, title, body) in enumerate(segments):
    x = .7 + i * 4.15
    rect(s, x, 2.55, 3.7, 2.7, INK_2, RGBColor(60, 105, 95))
    add_text(s, n, x + .25, 2.78, .4, .25, 10, LIME, True)
    add_text(s, title, x + .25, 3.3, 3.0, .38, 20, WHITE, True)
    add_text(s, body, x + .25, 3.98, 3.05, 1.05, 10, PALE)
add_text(s, 'VALIDATION NEEDED', .72, 5.86, 1.55, .22, 8, LIME, True)
add_text(s, 'Customer interviews → hiring volume → willingness to pay → sales-cycle evidence', 2.42, 5.82, 8.8, .35, 15, WHITE)
footer(s)

# 8 — Business model
s = new_slide()
header(s, 8, 'Business model', 'Land with a vacancy. Grow with the team.', 'Initial monetization is configurable; prices below are proposed hypotheses for customer testing—not approved live pricing.')
tiers = [
    ('POST', 'QAR 499', 'Published vacancy\nApplicant tracking\nTeam workspace'),
    ('MATCH', 'QAR 999', 'Everything in Post\nEvidence-led matching\nCandidate report'),
    ('INTERVIEW', 'QAR 1,499', 'Everything in Match\nText interview workflow\nAssessment report'),
]
for i, (name, price, body) in enumerate(tiers):
    x = .72 + i * 4.18
    fill = INK if i == 1 else CREAM
    c = WHITE if i == 1 else INK
    m = PALE if i == 1 else MUTED
    rect(s, x, 2.35, 3.78, 3.62, fill, fill if i == 1 else LINE)
    add_text(s, name, x + .3, 2.68, 1.4, .22, 9, LIME if i == 1 else MUTED, True)
    add_text(s, price, x + .3, 3.18, 2.9, .5, 25, c, True)
    add_text(s, 'per vacancy', x + .3, 3.78, 2.0, .25, 10, m)
    add_text(s, body, x + .3, 4.35, 2.9, 1.15, 12, m)
add_text(s, 'Future revenue: recruiter seats · subscriptions · credits · enterprise plans · integrations', .75, 6.39, 10.7, .3, 12, MUTED)
footer(s)

# 9 — Go to market
s = new_slide(CREAM)
header(s, 9, 'Go-to-market', 'Founder-led, proof-led, locally anchored.', 'Win trust through design partners and measurable improvements before scaling paid acquisition.')
gtm = [
    ('0–3 months', 'Discover', '25 employer interviews\n5 design partners\nBaseline hiring metrics'),
    ('3–6 months', 'Prove', 'Pilot real vacancies\nMeasure review time & completion\nPublish customer case studies'),
    ('6–12 months', 'Convert', 'Paid vacancy packages\nFounder-led outbound\nEcosystem partnerships'),
    ('12–18 months', 'Scale', 'Recruiter seats\nMid-market pipeline\nGCC market validation'),
]
for i, (period, title, body) in enumerate(gtm):
    x = .68 + i * 3.12
    add_text(s, period.upper(), x, 2.42, 2.5, .22, 9, MUTED, True)
    add_text(s, title, x, 2.92, 2.5, .36, 19, INK, True)
    line = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(3.55), Inches(2.65), Inches(.035))
    line.fill.solid()
    line.fill.fore_color.rgb = LIME
    line.line.fill.background()
    add_text(s, body, x, 3.9, 2.45, 1.3, 12, MUTED)
add_text(s, 'North-star pilot evidence', .7, 5.78, 2.5, .3, 15, INK, True)
add_text(s, 'Time-to-first-review  ·  shortlist quality  ·  candidate completion  ·  recruiter retention', 3.4, 5.79, 8.9, .35, 13, MUTED)
footer(s)

# 10 — Defensibility
s = new_slide()
header(s, 10, 'Defensibility', 'The moat compounds through accountable data.', 'Technology is necessary; trusted workflows, decision history and employer adoption create durable advantage.')
center_x, center_y = 5.2, 2.45
rect(s, center_x, center_y, 2.9, 1.3, INK, INK)
add_text(s, 'Hirewise\ntrust layer', center_x, center_y + .28, 2.9, .72, 20, WHITE, True, align=PP_ALIGN.CENTER)
moats = [
    (0.8, 2.25, 'Localized workflow', 'Qatar-first operations, payments and policies'),
    (9.3, 2.25, 'Evidence graph', 'Criteria → sources → scores → human outcomes'),
    (2.25, 4.65, 'Versioned history', 'Reproducible policies and accountable overrides'),
    (7.45, 4.65, 'Customer learning', 'Role templates, rubrics and workflow benchmarks'),
]
for x, y, title, body in moats:
    rect(s, x, y, 3.25, 1.35, CREAM)
    add_text(s, title, x + .2, y + .2, 2.8, .28, 14, INK, True)
    add_text(s, body, x + .2, y + .67, 2.8, .45, 10, MUTED)
add_text(s, 'PRIVACY + FAIRNESS + SECURITY', 4.28, 6.35, 4.8, .28, 11, INK, True, align=PP_ALIGN.CENTER)
footer(s)

# 11 — Progress
s = new_slide(INK)
header(s, 11, 'Execution status', 'Working foundation. Pre-production company.', 'The codebase demonstrates the workflow; commercial validation and production infrastructure are the next inflection points.', dark=True)
metrics = [('33', 'automated tests'), ('4', 'user experiences'), ('1', 'tenant-aware platform'), ('0', 'live customer claims')]
for i, (value, label) in enumerate(metrics):
    x = .75 + i * 3.08
    add_text(s, value, x, 2.42, 2.4, .62, 30, LIME, True)
    add_text(s, label, x, 3.13, 2.4, .32, 12, WHITE, True)
line_y = 4.06
items = [
    ('BUILT', 'Accounts · jobs · applications · CV quarantine · audit · text interview · payment sandbox'),
    ('VALIDATED', 'Authorization tests · mobile/desktop browser flow · PostgreSQL CI · container build'),
    ('OPEN', 'Customer discovery · semantic AI · live payments · production hosting · legal/security review'),
]
for i, (label, body) in enumerate(items):
    y = line_y + i * .67
    add_text(s, label, .75, y, 1.35, .22, 8, LIME, True)
    add_text(s, body, 2.2, y - .03, 10.1, .32, 12, PALE)
footer(s)

# 12 — Roadmap
s = new_slide()
header(s, 12, '18-month roadmap', 'Technical foundation → repeatable revenue.', 'Milestones are management targets and depend on funding, hiring, legal review and customer access.')
roadmap = [
    ('NOW', 'Foundation', 'Product workflow\nSecurity baseline\nInvestor/customer discovery'),
    ('0–6M', 'Pilot', 'Production hosting\n5 design partners\nReviewed matching v1'),
    ('6–12M', 'Commercialize', 'Qatar payments\nPaid vacancies\nCase-study evidence'),
    ('12–18M', 'Repeat', 'Recruiter seats\nMid-market pipeline\nGCC validation'),
]
for i, (period, title, body) in enumerate(roadmap):
    x = .73 + i * 3.08
    circle(s, x, 2.48, .45, INK if i == 0 else LIME, INK if i == 0 else LIME)
    if i < 3:
        shape = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x + .45), Inches(2.69), Inches(2.62), Inches(.025))
        shape.fill.solid()
        shape.fill.fore_color.rgb = LINE
        shape.line.fill.background()
    add_text(s, period, x, 3.22, 1.35, .22, 9, MUTED, True)
    add_text(s, title, x, 3.72, 2.6, .34, 18, INK, True)
    add_text(s, body, x, 4.32, 2.55, 1.15, 12, MUTED)
rect(s, .73, 5.92, 12.0, .55, PALE, PALE)
add_text(s, 'Every phase gated by human oversight, privacy review and measurable customer value.', 1.0, 6.08, 11.4, .25, 12, INK, True, align=PP_ALIGN.CENTER)
footer(s)

# 13 — Funding
s = new_slide(CREAM)
header(s, 13, 'Illustrative pre-seed', 'QAR 1.5M to reach a commercial proof point.', 'A management scenario for discussion—not a financing commitment, valuation or approved budget.')
allocations = [('45%', 'Product & engineering', INK), ('20%', 'Security, legal & compliance', BLUE), ('20%', 'Pilot delivery & sales', RGBColor(112, 143, 71)), ('10%', 'Infrastructure & AI', RGBColor(99, 82, 148)), ('5%', 'Contingency', MUTED)]
start_x = .75
for pct, label, color in allocations:
    width = float(pct[:-1]) / 100 * 11.85
    shape = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(start_x), Inches(2.5), Inches(width), Inches(.68))
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    start_x += width
for i, (pct, label, color) in enumerate(allocations):
    y = 3.62 + i * .48
    x = .77
    circle(s, x, y, .22, color, color)
    add_text(s, pct, x + .38, y - .02, .55, .25, 11, INK, True)
    add_text(s, label, x + 1.0, y - .02, 4.5, .3, 12, MUTED)
rect(s, 7.05, 3.68, 5.3, 2.05, WHITE, LINE)
add_text(s, 'Target proof point', 7.38, 3.98, 2.3, .3, 16, INK, True)
add_text(s, 'Production-ready Qatar launch\n5–10 design partners\nFirst paid customers + measured retention', 7.38, 4.55, 4.35, .87, 12, MUTED)
add_text(s, 'Assumed runway: ~18 months. Founder compensation, hiring plan and valuation remain to be agreed.', .78, 6.42, 11.7, .28, 10, MUTED)
footer(s)

# 14 — Team
s = new_slide()
header(s, 14, 'Team', 'The founding story belongs here.', 'Investor circulation should include verified founder names, roles, relevant experience, ownership and full-time commitment.')
rect(s, .73, 2.35, 7.55, 3.45, CREAM)
add_text(s, 'Founder / CEO', 1.08, 2.78, 2.3, .35, 20, INK, True)
add_text(s, '[Add verified name]', 1.08, 3.35, 3.0, .28, 14, MUTED)
add_rich_lines(s, ['Relevant industry insight', 'Commercial/customer access', 'Why this founder is uniquely positioned'], 1.08, 4.05, 5.5, 1.2, 12, MUTED, True, 7)
rect(s, 8.57, 2.35, 4.05, 3.45, INK, INK)
add_text(s, 'Near-term key hires', 8.92, 2.78, 3.2, .35, 18, WHITE, True)
add_rich_lines(s, ['Senior product engineer', 'Recruitment domain lead', 'Enterprise sales / partnerships', 'Fractional security & privacy counsel'], 8.92, 3.52, 3.15, 1.65, 12, PALE, True, 8)
add_text(s, 'Do not present placeholder biographies as facts.', .75, 6.35, 5.8, .25, 10, MUTED)
footer(s)

# 15 — Closing
s = new_slide(INK)
circle(s, 9.0, -1.2, 5.4, INK_2, INK_2)
circle(s, 10.15, .15, 3.1, LIME, LIME)
s.shapes.add_picture(str(LOGO), Inches(10.55), Inches(.55), width=Inches(2.0), height=Inches(2.0))
add_text(s, 'THE OPPORTUNITY', .82, .75, 3.0, .25, 9, LIME, True)
add_text(s, 'Build the trust layer\nfor better hiring.', .82, 1.45, 7.65, 1.45, 36, WHITE, True)
add_text(s, 'Hirewise is seeking design partners, recruitment expertise\nand pre-seed capital to move from a working foundation\nto a validated commercial platform.', .84, 3.47, 7.0, 1.25, 14, PALE)
rect(s, .82, 5.08, 7.0, .85, INK_2, INK_2)
add_text(s, 'SMARTER HIRING.', 1.15, 5.39, 2.1, .25, 11, WHITE, True)
add_text(s, 'FAIRER DECISIONS.', 3.45, 5.39, 2.3, .25, 11, LIME, True)
add_text(s, 'HUMAN CONTROL.', 5.95, 5.39, 1.55, .25, 11, WHITE, True)
add_text(s, 'Contact: [add founder email]  ·  Doha, Qatar', .84, 6.75, 6.4, .25, 10, PALE)

# 16 — Sources and assumptions
s = new_slide(WHITE)
header(s, 16, 'Appendix', 'Sources, assumptions & diligence notes', 'Prepared 13 September 2026. Verify all external and company data immediately before investor circulation.')
sources = [
    'Qatar MCIT — Digital Agenda 2030: https://www.mcit.gov.qa/en/nda',
    'MCIT Digital Agenda 2030 summary — ~QAR 40B cumulative impact and 26,000 ICT jobs by 2030: https://www.mcit.gov.qa/wp-content/uploads/sites/4/2024/09/digital_agenda_2030_summary_english.pdf',
    'Qatar Digital Economy Hub — ecosystem context: https://digitaleconomyhub.mcit.gov.qa/en',
    'Qatar Central Bank — retail payment systems and QPAY context: https://www.qcb.gov.qa/en/Pages/Retail-payment-systems.aspx',
    'Hirewise implementation evidence — repository tests, CI and docs/MILESTONES.md; no live customer/traction claim.',
]
add_rich_lines(s, sources, .75, 2.28, 11.85, 2.75, 8, MUTED, True, 6)
rect(s, .75, 5.23, 11.85, 1.05, CREAM)
add_text(s, 'Management assumptions', 1.02, 5.5, 2.4, .28, 13, INK, True)
add_text(s, 'Proposed pricing, QAR 1.5M funding ask, 18-month runway, design-partner targets and roadmap are illustrative. TAM/SAM/SOM, unit economics, founder details and commercial traction require primary validation.', 3.22, 5.42, 8.95, .55, 10, MUTED)
footer(s, 'CONFIDENTIAL · VERIFY BEFORE EXTERNAL CIRCULATION')

# Core metadata
prs.core_properties.title = 'Hirewise Investor Presentation'
prs.core_properties.subject = 'AI-assisted recruitment platform for Qatar'
prs.core_properties.author = 'Hirewise'
prs.core_properties.keywords = 'Hirewise, recruitment, Qatar, investor, startup, responsible AI'
prs.core_properties.comments = 'Figures labelled as assumptions require owner validation before circulation.'

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
prs.save(OUTPUT)
print(f'Created {OUTPUT} with {len(prs.slides)} slides.')
