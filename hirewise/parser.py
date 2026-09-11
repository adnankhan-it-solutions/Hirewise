"""Bounded subprocess parser; invoked only after successful malware scanning."""
import io
import json
import resource
import sys

from docx import Document
from pypdf import PdfReader


def extract(data, extension):
    if extension == '.pdf':
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted or len(reader.pages) > 60:
            raise ValueError('unsupported_pdf')
        sections = [{'section': f'Page {i + 1}', 'text': page.extract_text()[:20000]} for i, page in enumerate(reader.pages)]
    elif extension == '.docx':
        document = Document(io.BytesIO(data))
        sections = [{'section': f'Paragraph {i + 1}', 'text': p.text[:10000]} for i, p in enumerate(document.paragraphs[:2000]) if p.text.strip()]
        for table_index, table in enumerate(document.tables[:30]):
            for row_index, row in enumerate(table.rows[:100]):
                sections.append({'section': f'Table {table_index + 1}, row {row_index + 1}', 'text': ' | '.join(c.text for c in row.cells)[:10000]})
    elif extension == '.txt':
        sections = [{'section': f'Line {i + 1}', 'text': text[:10000]} for i, text in enumerate(data.decode('utf-8').splitlines()[:5000]) if text.strip()]
    else:
        raise ValueError('unsupported_extension')
    if sum(len(s['text']) for s in sections) > 300000:
        raise ValueError('extraction_limit')
    return sections


if __name__ == '__main__':
    resource.setrlimit(resource.RLIMIT_CPU, (20, 20))
    resource.setrlimit(resource.RLIMIT_AS, (512 * 1024 * 1024, 512 * 1024 * 1024))
    data = sys.stdin.buffer.read(5 * 1024 * 1024 + 1)
    if len(data) > 5 * 1024 * 1024:
        raise ValueError('input_limit')
    print(json.dumps(extract(data, sys.argv[1])))
